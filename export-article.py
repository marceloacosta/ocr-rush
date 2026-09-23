"""Regenerate the editable Word article from the standalone page.

Run: python3 export-article.py (requires python-docx and beautifulsoup4).
The site build uses the generated document and needs no Python dependencies.
"""
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parent / "public" / "article"
GAME_URL = "https://www.marcelops.com/ocr-rush/"
soup = BeautifulSoup((ROOT / "index.html").read_text(), "html.parser")
document = Document()
document.core_properties.title = soup.h1.get_text()
document.core_properties.author = "Marcelo Acosta"
document.core_properties.subject = "Companion article for OCR Rush"
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.8)
normal = document.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(10)
normal.paragraph_format.line_spacing = 1.15
for name, size in [("Title", 26), ("Subtitle", 13), ("Heading 1", 17)]:
    document.styles[name].font.name = "Calibri"
    document.styles[name].font.size = Pt(size)
    document.styles[name].font.color.rgb = RGBColor.from_string("20354A")


def append_inline(paragraph, node, bold=False, italic=False, parent=None):
    """Keep native Word text, emphasis and external hyperlinks editable."""
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
        if parent is not None:
            parent.append(run._r)
        return
    if node.name == "a":
        href = node.get("href", "")
        if href == "../":
            href = GAME_URL
        if not href.startswith("https://"):
            raise ValueError(f"Export link must be public HTTPS: {href}")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
        for child in node.children:
            append_inline(paragraph, child, bold, italic, hyperlink)
        for run in hyperlink.findall(qn("w:r")):
            properties = run.find(qn("w:rPr"))
            if properties is None:
                properties = OxmlElement("w:rPr")
                run.insert(0, properties)
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "0053B4")
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            properties.extend([color, underline])
        paragraph._p.append(hyperlink)
        return
    for child in node.children:
        append_inline(paragraph, child, bold or node.name == "strong", italic or node.name == "em", parent)


document.add_paragraph(soup.h1.get_text(), "Title")
document.add_paragraph(soup.select_one(".subtitle").get_text(), "Subtitle")
for element in soup.select_one(".article-body").find_all(recursive=False):
    if element.name == "h2":
        label = element.select_one(".level-label")
        if label:
            prefix = label.get_text() + ": "
            label.extract()
        else:
            prefix = ""
        document.add_heading(prefix + element.get_text(), level=1)
    elif element.name == "figure":
        picture = document.add_picture(str(ROOT / "document-journey.png"), width=Inches(6.3))
        properties = picture._inline.docPr
        properties.set("descr", "An uploaded PDF is prepared so its text and table regions can be located. An AI model then reads those regions and returns their contents. Work can wait in a queue before either stage.")
        document.paragraphs[-1].paragraph_format.keep_with_next = True
        caption = document.add_paragraph(element.figcaption.get_text())
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(10)
    elif element.name == "p":
        paragraph = document.add_paragraph()
        for child in element.children:
            append_inline(paragraph, child)
    else:
        raise ValueError(f"Unrecognized article element: {element.name}")

output = ROOT / "ocr-rush-companion.docx"
document.save(output)
print(f"Created {output}")
